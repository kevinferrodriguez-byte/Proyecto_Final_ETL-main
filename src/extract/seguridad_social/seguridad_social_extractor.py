from datetime import datetime, timezone
from io import BytesIO
import zipfile

from openpyxl import load_workbook

from src.extract.bronze_storage import BronzeStorage
from src.extract.http_client import HttpClient
from src.utils.run_manifest import RunManifest
from src.logger import get_logger

XLSX_SIGNATURE = b"PK\x03\x04"


class SeguridadSocialExtractor:
    def __init__(self, config):
        self.source_name = config["name"]
        self.files = config["files"]
        self.http = HttpClient(config["request"])
        self.storage = BronzeStorage(config["storage"]["payloads_path"])
        self.manifest = RunManifest(config["storage"]["manifests_path"])
        self.logger = get_logger(__name__)

    def extract_files(self, file_ids):
        started_at = datetime.now(timezone.utc)
        run = {"run_id": f"seguridad_social_{started_at.strftime('%Y%m%dT%H%M%S.%fZ')}", "started_at_utc": self._iso(started_at)}
        records = []
        try:
            for file_id in file_ids:
                records.append(self.extract_file(file_id))
        except (OSError, ValueError, RuntimeError) as error:
            self._write_manifest(run, file_ids, records, error)
            raise
        return self._write_manifest(run, file_ids, records, None)

    def extract_file(self, file_id):
        file = self.files.get(file_id)
        if file is None:
            raise ValueError(f"El archivo {file_id} no está configurado en sources.seguridad_social.files de config/config.yaml.")
        response = self.http.get(file["url"])
        validations = self._validate(file_id, file, response.content)
        stored = self.storage.save(response.content, file_id, file["extension"])
        record = {
            "source": self.source_name,
            "file_id": file_id,
            "name": file["name"],
            "publisher": file["publisher"],
            "url": response.url,
            "status_code": response.status_code,
            "content_type": response.headers.get("Content-Type"),
            "size_bytes": stored["size_bytes"],
            "sha256": stored["sha256"],
            "downloaded_at_utc": stored["downloaded_at_utc"],
            "path": stored["path"],
            "validations": validations,
        }
        size = f"{stored['size_bytes'] / 1_000:.1f} KB".replace(".", ",")
        self.logger.ok(f"{file_id} · {file['name']}: HTTP {response.status_code}, {size}, hoja '{file['sheet']}' presente")
        return record

    def close(self):
        self.http.close()

    @property
    def closed(self):
        """True cuando la conexión HTTP del extractor está cerrada."""
        return self.http.closed

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        self.close()
        return False

    def _validate(self, file_id, file, content):
        description = f"El archivo {file_id} de la Seguridad Social"
        if not content:
            raise ValueError(f"{description} llegó vacío; no se guarda ningún archivo.")
        if not content.startswith(XLSX_SIGNATURE):
            raise ValueError(f"{description} no es un libro Excel .{file['extension']} (¿la URL devolvió una página HTML?); no se guarda ningún archivo.")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True)
        except (zipfile.BadZipFile, OSError, KeyError, ValueError) as error:
            raise ValueError(f"{description} no se puede abrir como libro Excel: {error}; no se guarda ningún archivo.") from error
        sheets = list(workbook.sheetnames)
        workbook.close()
        required = [file["sheet"], *file.get("additional_sheets", [])]
        missing = [sheet for sheet in required if sheet not in sheets]
        if missing:
            raise ValueError(f"{description} no contiene las hojas {missing} (hojas: {', '.join(sheets)}); no se guarda ningún archivo.")
        return {"xlsx_valid": True, "sheet_present": True, "sheet": file["sheet"], "required_sheets": required, "sheets": sheets}

    def _write_manifest(self, run, file_ids, records, error):
        manifest = {
            "run_id": run["run_id"],
            "source": self.source_name,
            "status": "fallida" if error else "completada",
            "error": str(error) if error else None,
            "started_at_utc": run["started_at_utc"],
            "finished_at_utc": self._iso(datetime.now(timezone.utc)),
            "files_requested": list(file_ids),
            "payloads": [self._manifest_entry(record) for record in records],
        }
        return {"run_id": run["run_id"], "manifest_path": self.manifest.write(manifest), "records": records}

    def _manifest_entry(self, record):
        entry = {key: value for key, value in record.items() if key != "path"}
        entry["path"] = self.manifest.relative_path(record["path"])
        return entry

    def _iso(self, moment):
        return moment.isoformat(timespec="microseconds").replace("+00:00", "Z")
